"""Bank of England 'latest yield curve data' zip — the nominal (GLC Nominal)
gilt spot curve. No API key: this downloads a `/-/media/...` asset path,
which robots.txt allows (unlike the disallowed /boeapps/iadb CSV export
`fetchers/cycle.py` used to note as the only keyless daily source).

The zip holds one workbook per curve family for the current month; we read
the nominal workbook's "4. spot curve" sheet, whose header row is tenors in
years (0.5, 1, 1.5, ... 40) and whose data rows are one business day each.
"""
from __future__ import annotations

import io
import zipfile
from datetime import date, datetime

import openpyxl

from collector.http import GetBytes

ZIP_URL = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/latest-yield-curve-data.zip"
NOMINAL_HINT = "glc nominal"
SPOT_SHEET = "4. spot curve"


def _fmt_tenor(years: float) -> str:
    """Header float -> the label config.yaml's `boe:` keys use ('10' not
    '10.0'; '0.5' as-is)."""
    return str(int(years)) if float(years).is_integer() else str(years)


def _find_nominal_workbook(names: list[str]) -> str:
    for name in names:
        if NOMINAL_HINT in name.lower():
            return name
    raise ValueError(f"no nominal workbook ({NOMINAL_HINT!r}) found in {names}")


def parse_spot_curve(xlsx_bytes: bytes, tenors: list[str]) -> dict[str, list[tuple[date, float]]]:
    """{tenor label: [(date, pct), ...]} for the requested tenor labels, read
    off the '4. spot curve' sheet's 'years:' header row."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    if SPOT_SHEET not in wb.sheetnames:
        raise ValueError(f"{SPOT_SHEET!r} sheet not found (have {wb.sheetnames})")
    ws = wb[SPOT_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next((i for i, r in enumerate(rows) if r and r[0] == "years:"), None)
    if header_idx is None:
        raise ValueError(f"{SPOT_SHEET}: no 'years:' header row found")
    header = rows[header_idx]
    wanted = set(tenors)
    col_for: dict[str, int] = {}
    for col, val in enumerate(header):
        if col == 0 or not isinstance(val, (int, float)):
            continue
        label = _fmt_tenor(val)
        if label in wanted:
            col_for[label] = col
    missing = wanted - set(col_for)
    if missing:
        raise ValueError(f"{SPOT_SHEET}: tenor(s) not found: {sorted(missing)}")
    out: dict[str, list[tuple[date, float]]] = {t: [] for t in tenors}
    for r in rows[header_idx + 1:]:
        if not r or not isinstance(r[0], datetime):
            continue
        d = r[0].date()
        for tenor, col in col_for.items():
            if col < len(r) and isinstance(r[col], (int, float)):
                out[tenor].append((d, float(r[col])))
    return out


async def fetch_spot_curve(tenors: list[str], get_bytes: GetBytes) -> dict[str, list[tuple[date, float]]]:
    zip_bytes = await get_bytes(ZIP_URL)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        xlsx_bytes = zf.read(_find_nominal_workbook(zf.namelist()))
    return parse_spot_curve(xlsx_bytes, tenors)
